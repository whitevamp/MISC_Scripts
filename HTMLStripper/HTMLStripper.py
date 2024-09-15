import os
import tkinter as tk
from tkinter import filedialog, ttk
from bs4 import BeautifulSoup
import docx
import pandas as pd
from fpdf import FPDF
from ebooklib import epub
import language_tool_python
from spellchecker import SpellChecker
from tqdm import tqdm

# Initialize spell checker and grammar checker
spell = SpellChecker()
tool = language_tool_python.LanguageTool('en-US')

# Function to strip HTML and JavaScript
def strip_html_js(content):
    soup = BeautifulSoup(content, 'html.parser')
    for script in soup(["script", "style"]):
        script.decompose()
    return soup.get_text()

# Function to spell check content
def spell_check(content):
    misspelled = spell.unknown(content.split())
    for word in misspelled:
        correction = spell.correction(word)
        if correction is not None:
            content = content.replace(word, correction)
    return content

# Function to grammar check content
def grammar_check(content):
    matches = tool.check(content)
    return language_tool_python.utils.correct(content, matches)

# Function to format content (basic paragraphing)
def format_text(content):
    paragraphs = content.split('\n')
    formatted_content = '\n\n'.join([p.strip() for p in paragraphs if p.strip()])
    return formatted_content

# Function to save content in different formats
def save_content(content, output_directory, base_filename, output_format):
    print(f"Saving file: {base_filename} as {output_format}")  # Debugging line
    output_path = os.path.join(output_directory, base_filename + '.' + output_format)

    if output_format == 'txt':
        with open(output_path, 'w', encoding='utf-8') as file:
            file.write(content)
    elif output_format == 'doc':
        doc = docx.Document()
        doc.add_paragraph(content)
        doc.save(output_path)
    elif output_format == 'csv':
        df = pd.DataFrame([line.split() for line in content.split('\n')])
        df.to_csv(output_path, index=False)
    elif output_format == 'pdf':
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font("Arial", size=12)
        for line in content.split('\n'):
            try:
                pdf.cell(200, 10, txt=line.encode('latin-1', 'replace').decode('latin-1'), ln=True)
            except UnicodeEncodeError:
                pass
        pdf.output(output_path)
    elif output_format == 'epub':
        book = epub.EpubBook()
        book.set_identifier('id123456')
        book.set_title('Sample')
        book.set_language('en')
        chapter = epub.EpubHtml(title='Content', file_name='chap_01.xhtml', lang='en')
        chapter.content = f'<html><body>{content}</body></html>'
        book.add_item(chapter)
        book.spine = ['nav', chapter]
        epub.write_epub(output_path, book)
    elif output_format in ['html', 'htm']:
        with open(output_path, 'w', encoding='utf-8') as file:
            file.write(content)

# Recursive function to process HTML files
def process_directory(input_directory, output_directory, output_format, progress_ui):
    file_list = []
    for root, _, files in os.walk(input_directory):
        for file in files:
            if file.endswith(('.html', '.htm')):
                file_list.append(os.path.join(root, file))

    total_files = len(file_list)

    # Initialize progress in the UI and command line
    progress_ui['maximum'] = total_files
    progress_ui['value'] = 0

    # Progress bar for file processing in the command line
    for i, input_path in enumerate(tqdm(file_list, desc="Processing files", unit="file")):
        try:
            base_filename = os.path.splitext(os.path.basename(input_path))[0]
            with open(input_path, 'r', encoding='utf-8') as f:
                content = f.read()

            stripped_content = strip_html_js(content)
            checked_content = spell_check(stripped_content)
            grammar_corrected_content = grammar_check(checked_content)
            formatted_content = format_text(grammar_corrected_content)

            # If "All Formats" is selected, save in all formats
            if output_format == 'All Formats':
                formats = ['txt', 'doc', 'csv', 'pdf', 'epub', 'html', 'htm']
                for fmt in formats:
                    save_content(formatted_content, output_directory, base_filename, fmt)
            else:
                save_content(formatted_content, output_directory, base_filename, output_format)

            # Update UI progress bar
            progress_ui['value'] = i + 1
            root.update_idletasks()

        except Exception as e:
            print(f"Error processing {input_path}: {e}")

# Tkinter UI
def select_input_directory():
    directory = filedialog.askdirectory()
    if directory:
        input_directory_var.set(directory)

def select_output_directory():
    directory = filedialog.askdirectory()
    if directory:
        output_directory_var.set(directory)

def start_processing():
    input_directory = input_directory_var.get()
    output_directory = output_directory_var.get()
    output_format = format_var.get()
    if input_directory and output_directory and output_format:
        process_directory(input_directory, output_directory, output_format, progress_bar)
    else:
        print("Please select both input and output directories, and the output format.")

root = tk.Tk()
root.title("HTML Processor")

input_directory_var = tk.StringVar()
output_directory_var = tk.StringVar()
format_var = tk.StringVar()

tk.Label(root, text="Select Input Directory:").grid(row=0, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=input_directory_var, width=50).grid(row=0, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_input_directory).grid(row=0, column=2, padx=10, pady=10)

tk.Label(root, text="Select Output Directory:").grid(row=1, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=output_directory_var, width=50).grid(row=1, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_output_directory).grid(row=1, column=2, padx=10, pady=10)

tk.Label(root, text="Output Format:").grid(row=2, column=0, padx=10, pady=10)
format_options = ['txt', 'doc', 'csv', 'pdf', 'epub', 'html', 'htm', 'All Formats']
ttk.Combobox(root, textvariable=format_var, values=format_options).grid(row=2, column=1, padx=10, pady=10)

tk.Button(root, text="Start", command=start_processing).grid(row=3, column=0, columnspan=3, pady=20)

# Progress bar in UI
progress_bar = ttk.Progressbar(root, orient='horizontal', length=400, mode='determinate')
progress_bar.grid(row=4, column=0, columnspan=3, padx=10, pady=10)

root.mainloop()
